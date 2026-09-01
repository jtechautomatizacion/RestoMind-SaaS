"""
CU-05: Dashboard Financiero
Serie de ventas/gastos/ganancia por día + top platos, para alimentar
los gráficos del panel de administración ("cómo viaja mi dinero").

Las ventas se cuentan cuando la comanda pasa a estado 'cobrado'
(dinero realmente cobrado), no cuando se crea la comanda.
"""

from datetime import datetime, timedelta, date
from io import BytesIO
from typing import Optional, Tuple
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from backend.database import get_db
from backend.dependencies import get_cliente_id, get_tz_offset, get_usuario_actual
from backend.models import Comanda, ComandaPlato, Compra
from backend.schemas import DashboardResumen, DashboardSerieItem, DashboardTotales, TopPlatoItem, TopGastoItem
from backend.utils.security import validar_admin

router = APIRouter()

DIAS_SEMANA_ES = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def _calcular_rango(desde: Optional[str], hasta: Optional[str], dias: int, tz_offset: int) -> Tuple[date, date, datetime, datetime, timedelta]:
    """
    Las fechas que manda el usuario (y las que ve en el gráfico) son días del
    restaurante en su hora local. Los timestamps de la BD son UTC. Sin traducir
    entre ambos, una venta cobrada el viernes 22:33 en Lima se guarda como
    sábado 03:33 UTC y desaparece del filtro "hasta el viernes".
    Devuelve (inicio, fin, inicio_dt_utc, fin_dt_utc, desfase) — usado tanto por
    el resumen del dashboard como por el export a Excel, para no calcular el
    mismo rango dos veces de formas ligeramente distintas.
    """
    desfase = timedelta(minutes=tz_offset)

    if desde and hasta:
        inicio = datetime.fromisoformat(desde).date()
        fin = datetime.fromisoformat(hasta).date()
    else:
        hoy_local = (datetime.utcnow() - desfase).date()
        inicio = hoy_local - timedelta(days=dias - 1)
        fin = hoy_local

    # Extremos del rango local, expresados en UTC para comparar contra la BD.
    inicio_dt = datetime.combine(inicio, datetime.min.time()) + desfase
    fin_dt = datetime.combine(fin, datetime.max.time()) + desfase

    return inicio, fin, inicio_dt, fin_dt, desfase


@router.get("/dashboard/resumen", response_model=DashboardResumen)
def resumen_financiero(
    dias: int = Query(default=7, ge=1, le=30),
    desde: str = Query(default=None),
    hasta: str = Query(default=None),
    db: Session = Depends(get_db),
    cliente_id: str = Depends(get_cliente_id),
    usuario: str = Depends(get_usuario_actual),
    tz_offset: int = Depends(get_tz_offset),
):
    # Ganancias, gastos y quién atendió cada comanda son del dueño, no del
    # turno. Que el frontend muestre la pestaña Dinero solo al admin es UI,
    # no seguridad: sin esto, cualquier mozo autenticado leía las finanzas
    # completas del restaurante con un curl.
    validar_admin(db, usuario, cliente_id)

    inicio, fin, inicio_dt, fin_dt, desfase = _calcular_rango(desde, hasta, dias, tz_offset)

    comandas_cobradas = (
        db.query(Comanda)
        .filter(
            Comanda.cliente_id == cliente_id,
            Comanda.estado == "cobrado",
            Comanda.actualizado_en >= inicio_dt,
            Comanda.actualizado_en <= fin_dt,
        )
        .all()
    )

    compras_periodo = (
        db.query(Compra)
        .filter(
            Compra.cliente_id == cliente_id,
            Compra.estado == "registrado",
            Compra.fecha >= inicio.isoformat(),
            Compra.fecha <= fin.isoformat(),
        )
        .all()
    )

    ventas_por_dia: dict = {}
    comandas_por_dia: dict = {}
    for c in comandas_cobradas:
        # El timestamp es UTC; la barra del gráfico es el día local del restaurante.
        fecha_str = (c.actualizado_en - desfase).date().isoformat()
        ventas_por_dia[fecha_str] = ventas_por_dia.get(fecha_str, 0.0) + c.total_cuenta
        comandas_por_dia[fecha_str] = comandas_por_dia.get(fecha_str, 0) + 1

    gastos_por_dia: dict = {}
    for compra in compras_periodo:
        gastos_por_dia[compra.fecha] = gastos_por_dia.get(compra.fecha, 0.0) + compra.monto

    serie = []
    num_dias = (fin - inicio).days + 1
    for i in range(num_dias):
        fecha_str = (inicio + timedelta(days=i)).isoformat()
        ventas = round(ventas_por_dia.get(fecha_str, 0.0), 2)
        gastos = round(gastos_por_dia.get(fecha_str, 0.0), 2)
        serie.append(DashboardSerieItem(
            fecha=fecha_str,
            ventas=ventas,
            gastos=gastos,
            ganancia=round(ventas - gastos, 2),
            comandas=comandas_por_dia.get(fecha_str, 0),
        ))

    totales = DashboardTotales(
        ventas=round(sum(item.ventas for item in serie), 2),
        gastos=round(sum(item.gastos for item in serie), 2),
        ganancia=round(sum(item.ganancia for item in serie), 2),
        comandas=sum(item.comandas for item in serie),
    )

    top_platos_map: dict = {}
    if comandas_cobradas:
        ids_comandas = [c.id for c in comandas_cobradas]
        detalle = db.query(ComandaPlato).filter(ComandaPlato.comanda_id.in_(ids_comandas)).all()
        for cp in detalle:
            nombre = cp.plato.nombre if cp.plato else "Plato eliminado"
            acumulado = top_platos_map.setdefault(nombre, {"cantidad": 0, "revenue": 0.0})
            acumulado["cantidad"] += cp.cantidad
            acumulado["revenue"] += cp.subtotal

    top_platos = sorted(
        (TopPlatoItem(nombre=n, cantidad=v["cantidad"], ingresos=round(v["revenue"], 2)) for n, v in top_platos_map.items()),
        key=lambda item: item.ingresos,
        reverse=True,
    )[:5]

    top_gastos_map: dict = {}
    for compra in compras_periodo:
        categoria = compra.categoria or "Sin categoría"
        acumulado = top_gastos_map.setdefault(categoria, {"cantidad": 0, "monto": 0.0})
        acumulado["cantidad"] += 1
        acumulado["monto"] += compra.monto

    top_gastos = sorted(
        (TopGastoItem(categoria=c, cantidad=v["cantidad"], monto=round(v["monto"], 2)) for c, v in top_gastos_map.items()),
        key=lambda item: item.monto,
        reverse=True,
    )[:3]

    periodo_dias = num_dias if (desde and hasta) else dias
    return DashboardResumen(
        periodo_dias=periodo_dias, serie=serie, totales=totales, top_platos=top_platos, top_gastos=top_gastos
    )


_ENCABEZADO_FILL = PatternFill(start_color="FF5A3C", end_color="FF5A3C", fill_type="solid")
_ENCABEZADO_FONT = Font(color="FFFFFF", bold=True)
# Formato de celda de Excel, no un string armado a mano: así el número sigue
# siendo un número real (se puede sumar/graficar en Excel), solo que se
# *muestra* con el símbolo de soles y dos decimales, igual que en el sistema.
_FORMATO_SOLES = '"S/" #,##0.00'


def _escribir_hoja(ws, encabezados: list, filas: list, columnas_moneda: tuple = ()) -> None:
    ws.append(encabezados)
    for col in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col)
        celda.fill = _ENCABEZADO_FILL
        celda.font = _ENCABEZADO_FONT
        celda.alignment = Alignment(horizontal="center")

    for fila in filas:
        ws.append(fila)

    for fila_idx in range(2, len(filas) + 2):
        for col in columnas_moneda:
            ws.cell(row=fila_idx, column=col).number_format = _FORMATO_SOLES

    # Ancho automático simple: el más largo entre encabezado y valores de esa
    # columna, con un tope para que una descripción larga no vuelva la hoja
    # inmanejable en pantalla.
    for col in range(1, len(encabezados) + 1):
        letra = get_column_letter(col)
        largo = max([len(str(encabezados[col - 1]))] + [len(str(fila[col - 1])) for fila in filas]) if filas else len(str(encabezados[col - 1]))
        ws.column_dimensions[letra].width = min(largo + 3, 40)

    ws.freeze_panes = "A2"


@router.get("/dashboard/reporte-excel")
def reporte_excel(
    dias: int = Query(default=7, ge=1, le=30),
    desde: str = Query(default=None),
    hasta: str = Query(default=None),
    db: Session = Depends(get_db),
    cliente_id: str = Depends(get_cliente_id),
    usuario: str = Depends(get_usuario_actual),
    tz_offset: int = Depends(get_tz_offset),
):
    # Mismo motivo que /dashboard/resumen: el Excel lleva exactamente los
    # mismos datos financieros, en un archivo fácil de llevarse.
    validar_admin(db, usuario, cliente_id)

    inicio, fin, inicio_dt, fin_dt, desfase = _calcular_rango(desde, hasta, dias, tz_offset)

    comandas_cobradas = (
        db.query(Comanda)
        .filter(
            Comanda.cliente_id == cliente_id,
            Comanda.estado == "cobrado",
            Comanda.actualizado_en >= inicio_dt,
            Comanda.actualizado_en <= fin_dt,
        )
        .order_by(Comanda.actualizado_en)
        .all()
    )

    # A diferencia del resumen (que solo suma gastos 'registrado' porque un
    # 'cancelado' no debe afectar la ganancia), el Excel es un reporte de
    # auditoría: incluye ambos estados y una columna 'Estado' para que quien
    # lo revise vea también lo que se anuló, no solo lo que quedó vigente.
    compras_periodo = (
        db.query(Compra)
        .filter(
            Compra.cliente_id == cliente_id,
            Compra.fecha >= inicio.isoformat(),
            Compra.fecha <= fin.isoformat(),
        )
        .order_by(Compra.fecha)
        .all()
    )

    wb = Workbook()

    ws_ventas = wb.active
    ws_ventas.title = "Detalle de Ventas"
    filas_ventas = []
    for comanda in comandas_cobradas:
        fecha_local = comanda.actualizado_en - desfase
        for cp in comanda.comanda_platos:
            filas_ventas.append([
                fecha_local.strftime("%d/%m/%Y"),
                fecha_local.strftime("%H:%M"),
                comanda.id,
                comanda.numero_mesa,
                cp.plato.nombre if cp.plato else "Plato eliminado",
                cp.plato.categoria if cp.plato else "",
                cp.cantidad,
                cp.precio_unitario,
                cp.subtotal,
                comanda.total_cuenta,
                comanda.creado_por or "",
            ])
    _escribir_hoja(ws_ventas, [
        "Fecha", "Hora", "N° Comanda", "Mesa", "Plato", "Categoría",
        "Cantidad", "Precio unitario", "Subtotal", "Total de la comanda", "Atendido por",
    ], filas_ventas, columnas_moneda=(8, 9, 10))

    ws_gastos = wb.create_sheet("Detalle de Gastos")
    filas_gastos = [
        [
            datetime.fromisoformat(compra.fecha).strftime("%d/%m/%Y"),
            compra.descripcion,
            compra.categoria or "Sin categoría",
            compra.monto,
            "Registrado" if compra.estado == "registrado" else "Cancelado",
            compra.creado_por or "",
        ]
        for compra in compras_periodo
    ]
    _escribir_hoja(ws_gastos, [
        "Fecha", "Descripción", "Categoría", "Monto", "Estado", "Registrado por",
    ], filas_gastos, columnas_moneda=(4,))

    ventas_por_dia: dict = {}
    comandas_por_dia: dict = {}
    for c in comandas_cobradas:
        fecha_str = (c.actualizado_en - desfase).date().isoformat()
        ventas_por_dia[fecha_str] = ventas_por_dia.get(fecha_str, 0.0) + c.total_cuenta
        comandas_por_dia[fecha_str] = comandas_por_dia.get(fecha_str, 0) + 1

    gastos_por_dia: dict = {}
    for compra in compras_periodo:
        if compra.estado == "registrado":
            gastos_por_dia[compra.fecha] = gastos_por_dia.get(compra.fecha, 0.0) + compra.monto

    ws_resumen = wb.create_sheet("Resumen Diario")
    filas_resumen = []
    num_dias = (fin - inicio).days + 1
    for i in range(num_dias):
        fecha = inicio + timedelta(days=i)
        fecha_str = fecha.isoformat()
        ventas = round(ventas_por_dia.get(fecha_str, 0.0), 2)
        gastos = round(gastos_por_dia.get(fecha_str, 0.0), 2)
        filas_resumen.append([
            fecha.strftime("%d/%m/%Y"),
            DIAS_SEMANA_ES[fecha.weekday()],
            ventas,
            gastos,
            round(ventas - gastos, 2),
            comandas_por_dia.get(fecha_str, 0),
        ])
    _escribir_hoja(ws_resumen, [
        "Fecha", "Día", "Ventas", "Gastos", "Ganancia", "N° Comandas",
    ], filas_resumen, columnas_moneda=(3, 4, 5))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"reporte_{inicio.isoformat()}_a_{fin.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )
