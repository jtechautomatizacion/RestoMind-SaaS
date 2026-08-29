"""
Tests de integración para endpoints API.
Cubren el ciclo completo: Admin crea plato -> Mozo manda comanda ->
Cocina entrega -> Mozo cobra -> Dashboard refleja el dinero.
"""

import pytest


def test_health_endpoint(test_client):
    response = test_client.get('/health')
    assert response.status_code == 200
    assert response.json()['status'] == 'ok'


def test_respuestas_incluyen_cabeceras_de_seguridad(test_client):
    """CSP, X-Frame-Options, etc. deben ir en TODA respuesta, no solo en
    login — cualquier endpoint sirve para verificarlo."""
    response = test_client.get('/health')
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"
    assert "default-src 'self'" in response.headers["content-security-policy"]
    assert response.headers["referrer-policy"] == "strict-origin-when-cross-origin"


# ============ AUTENTICACIÓN ============

def test_login_correcto_devuelve_token(test_client_real_auth, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL, TEST_USUARIO_PASSWORD
    resp = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL, "password": TEST_USUARIO_PASSWORD,
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["access_token"]
    assert data["usuario"]["email"] == TEST_USUARIO_EMAIL
    assert data["usuario"]["rol"] == "admin"
    assert data["usuario"]["cliente_id"] == test_cliente.id


def test_login_password_incorrecta_rechaza(test_client_real_auth, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL
    resp = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL, "password": "password-equivocado",
    })
    assert resp.status_code == 401


def test_login_email_inexistente_rechaza(test_client_real_auth, test_cliente):
    resp = test_client_real_auth.post('/api/auth/login', json={
        "email": "no-existe@nadie.com", "password": "lo-que-sea",
    })
    assert resp.status_code == 401


def test_endpoint_protegido_sin_token_rechaza(test_client_real_auth, test_cliente):
    resp = test_client_real_auth.get('/api/platos')
    assert resp.status_code == 401


def test_endpoint_protegido_con_token_invalido_rechaza(test_client_real_auth, test_cliente):
    resp = test_client_real_auth.get('/api/platos', headers={"Authorization": "Bearer token-inventado"})
    assert resp.status_code == 401


def test_endpoint_protegido_con_token_valido_permite_acceso(test_client_real_auth, test_cliente, test_platos):
    from tests.conftest import TEST_USUARIO_EMAIL, TEST_USUARIO_PASSWORD
    login = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL, "password": TEST_USUARIO_PASSWORD,
    })
    token = login.json()["access_token"]

    resp = test_client_real_auth.get('/api/platos', headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 200
    assert len(resp.json()) == len(test_platos)


def test_auth_me_devuelve_usuario_del_token(test_client_real_auth, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL, TEST_USUARIO_PASSWORD
    login = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL, "password": TEST_USUARIO_PASSWORD,
    })
    token = login.json()["access_token"]

    resp = test_client_real_auth.get('/api/auth/me', headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 200
    assert resp.json()["email"] == TEST_USUARIO_EMAIL


def test_root_endpoint(test_client):
    # "/" redirige al frontend (no devuelve JSON): así el link de credenciales
    # por email lleva directo a la app, no a un endpoint de la API.
    response = test_client.get('/', follow_redirects=False)
    assert response.status_code in (302, 307)
    assert response.headers['location'] == '/static/index.html'


# ============ CU-01: PLATOS ============

def test_crear_y_listar_plato(test_client, test_cliente):
    payload = {"nombre": "Ceviche Clásico", "categoria": "Cebiches", "precio_venta": 45.0}
    resp = test_client.post('/api/platos', json=payload)
    assert resp.status_code == 201
    data = resp.json()
    assert data["nombre"] == "Ceviche Clásico"
    assert data["estado"] == "activo"
    assert data["cliente_id"] == "test-cliente-001"

    resp = test_client.get('/api/platos')
    assert resp.status_code == 200
    assert len(resp.json()) == 1


def test_crear_plato_precio_invalido_falla(test_client, test_cliente):
    resp = test_client.post('/api/platos', json={"nombre": "X", "categoria": "Y", "precio_venta": 0})
    assert resp.status_code == 422


def test_desactivar_plato_no_aparece_en_listado(test_client, test_platos):
    plato_id = test_platos[0].id
    resp = test_client.patch(f'/api/platos/{plato_id}/estado', json={"estado": "inactivo"})
    assert resp.status_code == 200
    assert resp.json()["estado"] == "inactivo"

    resp = test_client.get('/api/platos')
    nombres = [p["nombre"] for p in resp.json()]
    assert test_platos[0].nombre not in nombres


def test_eliminar_plato_sin_historial_lo_borra(test_client, test_platos):
    plato_id = test_platos[0].id
    resp = test_client.delete(f'/api/platos/{plato_id}')
    assert resp.status_code == 200
    assert resp.json()["eliminado"] is True

    resp = test_client.get('/api/platos?incluir_inactivos=true')
    ids = [p["id"] for p in resp.json()]
    assert plato_id not in ids


def test_eliminar_plato_con_historial_lo_archiva_sin_borrar(test_client, test_platos, test_mesas):
    plato_id = test_platos[0].id
    test_client.post('/api/comandas', json={
        "numero_mesa": 1,
        "platos": [{"plato_id": plato_id, "cantidad": 1}],
    })

    resp = test_client.delete(f'/api/platos/{plato_id}')
    assert resp.status_code == 200
    assert resp.json()["eliminado"] is False

    # Sigue existiendo (no se perdió el historial de la comanda), pero
    # desapareció de la carta activa.
    resp = test_client.get('/api/platos?incluir_inactivos=true')
    plato = next(p for p in resp.json() if p["id"] == plato_id)
    assert plato["estado"] == "inactivo"

    resp = test_client.get('/api/platos')
    ids = [p["id"] for p in resp.json()]
    assert plato_id not in ids


# ============ CU-02: COMANDAS ============

def test_crear_comanda_calcula_total_y_ocupa_mesa(test_client, test_platos, test_mesas):
    payload = {
        "numero_mesa": 1,
        "platos": [
            {"plato_id": test_platos[0].id, "cantidad": 2},  # 45 * 2 = 90
            {"plato_id": test_platos[1].id, "cantidad": 1},  # 5 * 1 = 5
        ],
    }
    resp = test_client.post('/api/comandas', json=payload)
    assert resp.status_code == 201
    data = resp.json()
    assert data["total_cuenta"] == 95.0
    assert data["estado"] == "cocina"
    assert len(data["platos"]) == 2

    mesas = test_client.get('/api/mesas').json()
    mesa1 = next(m for m in mesas if m["numero"] == 1)
    assert mesa1["estado"] == "ocupada"
    assert mesa1["cuenta_actual"] == 95.0


def test_crear_comanda_mesa_inexistente_falla(test_client, test_platos):
    payload = {"numero_mesa": 999, "platos": [{"plato_id": test_platos[0].id, "cantidad": 1}]}
    resp = test_client.post('/api/comandas', json=payload)
    assert resp.status_code == 404


def test_crear_comanda_plato_duplicado_suma_cantidades(test_client, test_platos, test_mesas):
    payload = {
        "numero_mesa": 1,
        "platos": [
            {"plato_id": test_platos[0].id, "cantidad": 1},
            {"plato_id": test_platos[0].id, "cantidad": 2},
        ],
    }
    resp = test_client.post('/api/comandas', json=payload)
    assert resp.status_code == 201
    data = resp.json()
    assert len(data["platos"]) == 1
    assert data["platos"][0]["cantidad"] == 3
    assert data["total_cuenta"] == 135.0


def test_crear_comanda_sin_platos_falla_validacion(test_client, test_mesas):
    resp = test_client.post('/api/comandas', json={"numero_mesa": 1, "platos": []})
    assert resp.status_code == 422


# ============ CU-03: MONITOR DE COCINA ============

def test_monitor_cocina_y_marcar_entregado(test_client, test_platos, test_mesas):
    payload = {"numero_mesa": 2, "platos": [{"plato_id": test_platos[0].id, "cantidad": 1}]}
    comanda = test_client.post('/api/comandas', json=payload).json()

    resp = test_client.get('/api/monitor/cocina')
    assert resp.status_code == 200
    ids = [c["id"] for c in resp.json()]
    assert comanda["id"] in ids

    resp = test_client.patch(f'/api/comandas/{comanda["id"]}/estado', json={"estado": "entregado"})
    assert resp.status_code == 200
    assert resp.json()["estado"] == "entregado"

    # Ya no debe aparecer en cocina
    resp = test_client.get('/api/monitor/cocina')
    ids = [c["id"] for c in resp.json()]
    assert comanda["id"] not in ids


def test_transicion_estado_invalida_falla(test_client, test_platos, test_mesas):
    payload = {"numero_mesa": 1, "platos": [{"plato_id": test_platos[0].id, "cantidad": 1}]}
    comanda = test_client.post('/api/comandas', json=payload).json()

    # No se puede pasar de 'cocina' directo a 'cobrado'
    resp = test_client.patch(f'/api/comandas/{comanda["id"]}/estado', json={"estado": "cobrado"})
    assert resp.status_code == 400


# ============ COBRO DE MESA (cierre del ciclo) ============

def test_no_se_puede_cobrar_con_platos_en_cocina(test_client, test_platos, test_mesas):
    payload = {"numero_mesa": 1, "platos": [{"plato_id": test_platos[0].id, "cantidad": 1}]}
    test_client.post('/api/comandas', json=payload)

    mesas = test_client.get('/api/mesas').json()
    mesa1_id = next(m for m in mesas if m["numero"] == 1)["id"]

    resp = test_client.post(f'/api/mesas/{mesa1_id}/cobrar')
    assert resp.status_code == 400


def test_cobrar_mesa_libera_mesa_y_suma_total(test_client, test_platos, test_mesas):
    payload = {"numero_mesa": 1, "platos": [{"plato_id": test_platos[0].id, "cantidad": 2}]}
    comanda = test_client.post('/api/comandas', json=payload).json()
    test_client.patch(f'/api/comandas/{comanda["id"]}/estado', json={"estado": "entregado"})

    mesas = test_client.get('/api/mesas').json()
    mesa1_id = next(m for m in mesas if m["numero"] == 1)["id"]

    resp = test_client.post(f'/api/mesas/{mesa1_id}/cobrar')
    assert resp.status_code == 200
    data = resp.json()
    assert data["total_cobrado"] == 90.0
    assert data["comandas_cerradas"] == 1

    mesas = test_client.get('/api/mesas').json()
    mesa1 = next(m for m in mesas if m["numero"] == 1)
    assert mesa1["estado"] == "disponible"
    assert mesa1["cuenta_actual"] == 0.0


# ============ CU-04: COMPRAS ============

def test_crear_compra_y_editar_mismo_dia(test_client, test_cliente):
    from datetime import datetime
    payload = {"descripcion": "Pescado rojo", "categoria": "Insumos", "monto": 85.5, "fecha": datetime.utcnow().date().isoformat()}
    resp = test_client.post('/api/compras', json=payload)
    assert resp.status_code == 201
    compra_id = resp.json()["id"]

    resp = test_client.patch(f'/api/compras/{compra_id}', json={"monto": 90.0})
    assert resp.status_code == 200
    assert resp.json()["monto"] == 90.0


def test_crear_compra_fecha_futura_falla(test_client, test_cliente):
    payload = {"descripcion": "Insumo", "monto": 10.0, "fecha": "2099-01-01"}
    resp = test_client.post('/api/compras', json=payload)
    assert resp.status_code == 400


def test_cancelar_compra(test_client, test_cliente):
    from datetime import datetime
    payload = {"descripcion": "Verduras", "monto": 30.0, "fecha": datetime.utcnow().date().isoformat()}
    compra_id = test_client.post('/api/compras', json=payload).json()["id"]

    resp = test_client.patch(f'/api/compras/{compra_id}/estado', json={"estado": "cancelado"})
    assert resp.status_code == 200
    assert resp.json()["estado"] == "cancelado"


def test_eliminar_compra_mismo_dia(test_client, test_cliente):
    from datetime import datetime
    payload = {"descripcion": "Hielo", "monto": 15.0, "fecha": datetime.utcnow().date().isoformat()}
    compra_id = test_client.post('/api/compras', json=payload).json()["id"]

    resp = test_client.delete(f'/api/compras/{compra_id}')
    assert resp.status_code == 204

    resp = test_client.get('/api/compras')
    ids = [c["id"] for c in resp.json()]
    assert compra_id not in ids


def test_eliminar_compra_dia_anterior_falla(test_client, test_cliente, test_db):
    from backend.models import Compra
    compra = Compra(cliente_id=test_cliente.id, descripcion="Gas", monto=50.0, fecha="2020-01-01", estado="registrado")
    test_db.add(compra)
    test_db.commit()

    resp = test_client.delete(f'/api/compras/{compra.id}')
    assert resp.status_code == 403


# ============ CU-05: DASHBOARD ============

def test_dashboard_refleja_venta_cobrada(test_client, test_platos, test_mesas):
    payload = {"numero_mesa": 1, "platos": [{"plato_id": test_platos[0].id, "cantidad": 1}]}
    comanda = test_client.post('/api/comandas', json=payload).json()
    test_client.patch(f'/api/comandas/{comanda["id"]}/estado', json={"estado": "entregado"})

    mesas = test_client.get('/api/mesas').json()
    mesa1_id = next(m for m in mesas if m["numero"] == 1)["id"]
    test_client.post(f'/api/mesas/{mesa1_id}/cobrar')

    resp = test_client.get('/api/dashboard/resumen?dias=7')
    assert resp.status_code == 200
    data = resp.json()
    assert data["totales"]["ventas"] == 45.0
    assert data["totales"]["comandas"] == 1
    assert len(data["top_platos"]) == 1
    assert data["top_platos"][0]["nombre"] == "Ceviche Clásico"


def test_dashboard_cuenta_venta_nocturna_en_el_dia_local(test_client, test_db, test_platos, test_mesas):
    """Una venta cobrada el viernes 22:33 en Lima se guarda como sábado 03:33 UTC.
    Debe contar el viernes (día del restaurante), no el sábado."""
    from datetime import datetime
    from backend.models import Comanda

    payload = {"numero_mesa": 1, "platos": [{"plato_id": test_platos[0].id, "cantidad": 1}]}
    comanda_id = test_client.post('/api/comandas', json=payload).json()["id"]
    test_client.patch(f'/api/comandas/{comanda_id}/estado', json={"estado": "entregado"})

    mesas = test_client.get('/api/mesas').json()
    mesa1_id = next(m for m in mesas if m["numero"] == 1)["id"]
    test_client.post(f'/api/mesas/{mesa1_id}/cobrar')

    comanda = test_db.query(Comanda).filter(Comanda.id == comanda_id).first()
    comanda.actualizado_en = datetime(2026, 8, 29, 3, 33)  # UTC
    test_db.commit()

    # Perú = UTC-5 → getTimezoneOffset() devuelve 300
    headers = {"X-TZ-Offset": "300"}
    resp = test_client.get('/api/dashboard/resumen?desde=2026-08-22&hasta=2026-08-28', headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["totales"]["ventas"] == 45.0
    dia_viernes = next(d for d in data["serie"] if d["fecha"] == "2026-08-28")
    assert dia_viernes["ventas"] == 45.0

    # Sin el header (servidor en UTC) la misma venta cae fuera del rango.
    sin_tz = test_client.get('/api/dashboard/resumen?desde=2026-08-22&hasta=2026-08-28').json()
    assert sin_tz["totales"]["ventas"] == 0.0


def test_reporte_excel_genera_las_3_hojas(test_client, test_platos, test_mesas):
    from io import BytesIO
    from openpyxl import load_workbook

    payload = {"numero_mesa": 1, "platos": [{"plato_id": test_platos[0].id, "cantidad": 2}]}
    comanda = test_client.post('/api/comandas', json=payload).json()
    test_client.patch(f'/api/comandas/{comanda["id"]}/estado', json={"estado": "entregado"})
    mesas = test_client.get('/api/mesas').json()
    mesa1_id = next(m for m in mesas if m["numero"] == 1)["id"]
    test_client.post(f'/api/mesas/{mesa1_id}/cobrar')

    from datetime import datetime
    test_client.post('/api/compras', json={
        "descripcion": "Insumos", "categoria": "Insumos", "monto": 50.0,
        "fecha": datetime.utcnow().date().isoformat(),
    })

    resp = test_client.get('/api/dashboard/reporte-excel?dias=7')
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Detalle de Ventas", "Detalle de Gastos", "Resumen Diario"]

    ws_ventas = wb["Detalle de Ventas"]
    assert ws_ventas.cell(row=1, column=1).value == "Fecha"
    assert ws_ventas.cell(row=2, column=5).value == "Ceviche Clásico"
    assert ws_ventas.cell(row=2, column=7).value == 2  # cantidad

    ws_gastos = wb["Detalle de Gastos"]
    assert ws_gastos.cell(row=2, column=2).value == "Insumos"

    ws_resumen = wb["Resumen Diario"]
    assert ws_resumen.cell(row=1, column=1).value == "Fecha"
    assert len(list(ws_resumen.iter_rows(min_row=2))) == 7


def test_dashboard_sin_ventas_no_falla(test_client, test_cliente):
    resp = test_client.get('/api/dashboard/resumen')
    assert resp.status_code == 200
    data = resp.json()
    assert data["totales"]["ventas"] == 0.0
    assert len(data["serie"]) == 7
    assert data["top_gastos"] == []


def test_dashboard_top_gastos_agrupa_por_categoria(test_client, test_cliente):
    from datetime import datetime
    hoy = datetime.utcnow().date().isoformat()
    test_client.post('/api/compras', json={"descripcion": "Pescado", "categoria": "Insumos", "monto": 100.0, "fecha": hoy})
    test_client.post('/api/compras', json={"descripcion": "Verduras", "categoria": "Insumos", "monto": 50.0, "fecha": hoy})
    test_client.post('/api/compras', json={"descripcion": "Luz", "categoria": "Servicios", "monto": 30.0, "fecha": hoy})

    resp = test_client.get('/api/dashboard/resumen?dias=7')
    assert resp.status_code == 200
    top_gastos = resp.json()["top_gastos"]
    assert top_gastos[0]["categoria"] == "Insumos"
    assert top_gastos[0]["monto"] == 150.0
    assert top_gastos[0]["cantidad"] == 2
    assert top_gastos[1]["categoria"] == "Servicios"
    assert top_gastos[1]["monto"] == 30.0


# ============ FOTOS DE PLATOS ============

# PNG 1x1 real (encabezado válido), para que la detección por firma binaria
# lo acepte igual que aceptaría una foto real tomada por un mozo.
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753"
    "de0000000c4944415478da6360606060000000050001a5f645400000000049454e44ae426082"
)


@pytest.fixture(autouse=True)
def _carpeta_imagenes_aislada(tmp_path, monkeypatch):
    """Redirige los uploads de imagen a una carpeta temporal por test,
    para no escribir archivos reales dentro del repo al correr la suite."""
    from backend.routes import platos as platos_module
    monkeypatch.setattr(platos_module, "CARPETA_IMAGENES", tmp_path / "platos")


def _crear_plato(test_client):
    payload = {"nombre": "Ceviche Clásico", "categoria": "Cebiches", "precio_venta": 45.0}
    return test_client.post('/api/platos', json=payload).json()["id"]


def test_subir_imagen_plato_valida(test_client, test_cliente):
    plato_id = _crear_plato(test_client)

    resp = test_client.post(
        f'/api/platos/{plato_id}/imagen',
        files={"archivo": ("foto.png", PNG_1PX, "image/png")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["imagen_url"] == f"/static/assets/platos/{plato_id}.png"

    # Debe reflejarse también al releer el plato
    resp = test_client.get('/api/platos')
    assert resp.json()[0]["imagen_url"] == f"/static/assets/platos/{plato_id}.png"


def test_subir_imagen_con_extension_falsa_falla(test_client, test_cliente):
    """El Content-Type que manda el cliente no es de fiar: se detecta por
    firma binaria real. Un .html disfrazado de imagen debe rechazarse."""
    plato_id = _crear_plato(test_client)

    resp = test_client.post(
        f'/api/platos/{plato_id}/imagen',
        files={"archivo": ("foto.png", b"<script>alert(1)</script>", "image/png")},
    )
    assert resp.status_code == 400


def test_subir_imagen_muy_pesada_falla(test_client, test_cliente):
    plato_id = _crear_plato(test_client)
    contenido_grande = PNG_1PX[:8] + b"\x00" * (5 * 1024 * 1024 + 1)

    resp = test_client.post(
        f'/api/platos/{plato_id}/imagen',
        files={"archivo": ("foto.png", contenido_grande, "image/png")},
    )
    assert resp.status_code == 400


def test_subir_imagen_plato_inexistente_falla(test_client, test_cliente):
    resp = test_client.post(
        '/api/platos/9999/imagen',
        files={"archivo": ("foto.png", PNG_1PX, "image/png")},
    )
    assert resp.status_code == 404


def test_reemplazar_imagen_no_deja_huerfanos(test_client, test_cliente, tmp_path):
    from backend.routes import platos as platos_module
    plato_id = _crear_plato(test_client)

    test_client.post(f'/api/platos/{plato_id}/imagen', files={"archivo": ("a.png", PNG_1PX, "image/png")})
    jpeg_1px = bytes.fromhex("ffd8ffe000104a46494600010100000100010000ffd9")
    resp = test_client.post(f'/api/platos/{plato_id}/imagen', files={"archivo": ("b.jpg", jpeg_1px, "image/jpeg")})

    assert resp.status_code == 200
    assert resp.json()["imagen_url"] == f"/static/assets/platos/{plato_id}.jpg"
    archivos = list(platos_module.CARPETA_IMAGENES.glob(f"{plato_id}.*"))
    assert len(archivos) == 1  # el .png viejo se borró, no quedó huérfano


def test_eliminar_imagen_plato(test_client, test_cliente):
    plato_id = _crear_plato(test_client)
    test_client.post(f'/api/platos/{plato_id}/imagen', files={"archivo": ("a.png", PNG_1PX, "image/png")})

    resp = test_client.delete(f'/api/platos/{plato_id}/imagen')
    assert resp.status_code == 200
    assert resp.json()["imagen_url"] is None


# ============ GESTIÓN DE MESAS (Admin) ============

def test_editar_mesa(test_client, test_mesas):
    mesa_id = test_mesas[0].id
    resp = test_client.patch(f'/api/mesas/{mesa_id}', json={"capacidad": 6, "ubicacion": "Patio"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["capacidad"] == 6
    assert data["ubicacion"] == "Patio"
    assert data["numero"] == test_mesas[0].numero  # no tocado, sigue igual


def test_editar_mesa_numero_duplicado_falla(test_client, test_mesas):
    mesa_id = test_mesas[0].id
    otro_numero = test_mesas[1].numero
    resp = test_client.patch(f'/api/mesas/{mesa_id}', json={"numero": otro_numero})
    assert resp.status_code == 400


def test_eliminar_mesa_disponible(test_client, test_mesas):
    mesa_id = test_mesas[0].id
    resp = test_client.delete(f'/api/mesas/{mesa_id}')
    assert resp.status_code == 204

    resp = test_client.get('/api/mesas')
    numeros = [m["numero"] for m in resp.json()]
    assert test_mesas[0].numero not in numeros


def test_eliminar_mesa_ocupada_falla(test_client, test_platos, test_mesas):
    mesa = test_mesas[0]
    test_client.post('/api/comandas', json={
        "numero_mesa": mesa.numero,
        "platos": [{"plato_id": test_platos[0].id, "cantidad": 1}],
    })

    resp = test_client.delete(f'/api/mesas/{mesa.id}')
    assert resp.status_code == 400


# ============ CATEGORÍAS DE LA CARTA (Admin) ============

def test_crear_y_listar_categoria(test_client, test_cliente):
    resp = test_client.post('/api/categorias', json={"nombre": "Cebiches"})
    assert resp.status_code == 201
    assert resp.json()["nombre"] == "Cebiches"

    resp = test_client.get('/api/categorias')
    assert resp.status_code == 200
    assert [c["nombre"] for c in resp.json()] == ["Cebiches"]


def test_crear_categoria_duplicada_falla(test_client, test_cliente):
    test_client.post('/api/categorias', json={"nombre": "Bebidas"})
    resp = test_client.post('/api/categorias', json={"nombre": "Bebidas"})
    assert resp.status_code == 400


def test_eliminar_categoria_sin_platos(test_client, test_cliente):
    categoria_id = test_client.post('/api/categorias', json={"nombre": "Postres"}).json()["id"]
    resp = test_client.delete(f'/api/categorias/{categoria_id}')
    assert resp.status_code == 204

    resp = test_client.get('/api/categorias')
    assert resp.json() == []


def test_eliminar_categoria_en_uso_falla(test_client, test_cliente):
    test_client.post('/api/categorias', json={"nombre": "Cebiches"})
    test_client.post('/api/platos', json={"nombre": "Ceviche Clásico", "categoria": "Cebiches", "precio_venta": 45.0})

    categorias = test_client.get('/api/categorias').json()
    categoria_id = categorias[0]["id"]

    resp = test_client.delete(f'/api/categorias/{categoria_id}')
    assert resp.status_code == 400


# ============ SUPERADMIN (panel del revendedor) ============

def _login_superadmin(client):
    from tests.conftest import TEST_SUPERADMIN_EMAIL, TEST_SUPERADMIN_PASSWORD
    resp = client.post('/api/superadmin/login', json={
        "email": TEST_SUPERADMIN_EMAIL, "password": TEST_SUPERADMIN_PASSWORD,
    })
    return resp.json()["access_token"]


def test_superadmin_login_correcto(test_client_real_auth, test_superadmin):
    from tests.conftest import TEST_SUPERADMIN_EMAIL, TEST_SUPERADMIN_PASSWORD
    resp = test_client_real_auth.post('/api/superadmin/login', json={
        "email": TEST_SUPERADMIN_EMAIL, "password": TEST_SUPERADMIN_PASSWORD,
    })
    assert resp.status_code == 200
    assert resp.json()["access_token"]
    assert resp.json()["email"] == TEST_SUPERADMIN_EMAIL


def test_superadmin_login_password_incorrecta(test_client_real_auth, test_superadmin):
    from tests.conftest import TEST_SUPERADMIN_EMAIL
    resp = test_client_real_auth.post('/api/superadmin/login', json={
        "email": TEST_SUPERADMIN_EMAIL, "password": "equivocada",
    })
    assert resp.status_code == 401


def test_superadmin_lista_clientes_con_stats(test_client_real_auth, test_superadmin, test_cliente, test_platos, test_mesas):
    token = _login_superadmin(test_client_real_auth)
    resp = test_client_real_auth.get('/api/superadmin/clientes', headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 200
    clientes = resp.json()
    assert len(clientes) == 1
    assert clientes[0]["id"] == test_cliente.id
    assert clientes[0]["num_platos"] == len(test_platos)
    assert clientes[0]["num_mesas"] == len(test_mesas)


def test_superadmin_crea_restaurante_nuevo(test_client_real_auth, test_superadmin):
    token = _login_superadmin(test_client_real_auth)
    resp = test_client_real_auth.post('/api/superadmin/clientes', headers={"Authorization": f"Bearer {token}"}, json={
        "nombre": "Pollería El Buen Sabor",
        "email": "contacto@buensabor.pe",
        "num_mesas": 5,
        "admin_nombre": "Carlos",
        "admin_email": "carlos@buensabor.pe",
        "admin_password": "clave123456",
    })
    assert resp.status_code == 201
    data = resp.json()
    assert data["id"] == "polleria-el-buen-sabor"
    assert data["num_mesas"] == 5

    # El nuevo admin ya puede loguearse en su propio restaurante.
    login = test_client_real_auth.post('/api/auth/login', json={
        "email": "carlos@buensabor.pe", "password": "clave123456",
    })
    assert login.status_code == 200
    assert login.json()["usuario"]["cliente_id"] == "polleria-el-buen-sabor"


def test_superadmin_suspender_cliente_bloquea_su_login(test_client_real_auth, test_superadmin, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL, TEST_USUARIO_PASSWORD
    token = _login_superadmin(test_client_real_auth)

    resp = test_client_real_auth.patch(
        f'/api/superadmin/clientes/{test_cliente.id}/estado',
        headers={"Authorization": f"Bearer {token}"}, json={"estado": "suspendido"},
    )
    assert resp.status_code == 200

    login = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL, "password": TEST_USUARIO_PASSWORD,
    })
    assert login.status_code == 403


def test_superadmin_resetea_password_de_admin(test_client_real_auth, test_superadmin, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL
    token = _login_superadmin(test_client_real_auth)

    resp = test_client_real_auth.patch(
        f'/api/superadmin/clientes/{test_cliente.id}/reset-password',
        headers={"Authorization": f"Bearer {token}"}, json={"nueva_password": "nuevaclave999"},
    )
    assert resp.status_code == 200

    login = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL, "password": "nuevaclave999",
    })
    assert login.status_code == 200


def test_token_de_restaurante_no_sirve_en_superadmin(test_client_real_auth, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL, TEST_USUARIO_PASSWORD
    login = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL, "password": TEST_USUARIO_PASSWORD,
    })
    token = login.json()["access_token"]

    resp = test_client_real_auth.get('/api/superadmin/clientes', headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 403


def test_token_de_superadmin_no_sirve_en_rutas_de_restaurante(test_client_real_auth, test_superadmin):
    token = _login_superadmin(test_client_real_auth)
    resp = test_client_real_auth.get('/api/platos', headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 403


def test_token_sin_tipo_explicito_no_accede_a_rutas_de_restaurante(test_client_real_auth, test_cliente):
    """get_cliente_id exige tipo == "usuario" explícito, no solo "no es
    superadmin" — un token que no declare tipo (uno viejo, o de un tipo
    futuro desconocido) no debe colarse solo por descarte."""
    import jwt
    from datetime import datetime, timedelta, timezone
    from backend.auth import ALGORITMO
    from backend.config import settings
    from tests.conftest import TEST_CLIENTE_ID

    payload = {
        "sub": "alguien@test.com",
        "cliente_id": TEST_CLIENTE_ID,
        "exp": datetime.now(timezone.utc) + timedelta(hours=1),
    }
    token = jwt.encode(payload, settings.secret_key, algorithm=ALGORITMO)

    resp = test_client_real_auth.get('/api/platos', headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 403


# ============ PERSONAL (usuarios de un restaurante) ============

def test_crear_y_listar_usuario_personal(test_client, test_cliente):
    resp = test_client.post('/api/usuarios', json={
        "nombre": "Pedro Mozo", "email": "pedro@test-restaurant.com", "password": "clave123", "rol": "mozo",
    })
    assert resp.status_code == 201
    assert resp.json()["rol"] == "mozo"

    resp = test_client.get('/api/usuarios')
    emails = [u["email"] for u in resp.json()]
    assert "pedro@test-restaurant.com" in emails


def test_nuevo_mozo_puede_loguearse_con_su_propio_rol(test_client_real_auth, test_cliente):
    from tests.conftest import TEST_CLIENTE_ID
    token_admin = _login_restaurante(test_client_real_auth)
    test_client_real_auth.post('/api/usuarios', json={
        "nombre": "Pedro Mozo", "email": "pedro2@test-restaurant.com", "password": "clave123456", "rol": "mozo",
    }, headers={"Authorization": f"Bearer {token_admin}"})

    login = test_client_real_auth.post('/api/auth/login', json={
        "email": "pedro2@test-restaurant.com", "password": "clave123456",
    })
    assert login.status_code == 200
    assert login.json()["usuario"]["rol"] == "mozo"
    assert login.json()["usuario"]["cliente_id"] == TEST_CLIENTE_ID


def test_crear_usuario_email_duplicado_falla(test_client, test_cliente):
    test_client.post('/api/usuarios', json={
        "nombre": "A", "email": "dup@test.com", "password": "clave123", "rol": "mozo",
    })
    resp = test_client.post('/api/usuarios', json={
        "nombre": "B", "email": "dup@test.com", "password": "clave123", "rol": "cajero",
    })
    assert resp.status_code == 400


def test_admin_no_puede_desactivarse_a_si_mismo(test_client, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL
    usuarios = test_client.get('/api/usuarios').json()
    mi_id = next(u["id"] for u in usuarios if u["email"] == TEST_USUARIO_EMAIL)

    resp = test_client.patch(f'/api/usuarios/{mi_id}', json={"estado": "inactivo"})
    assert resp.status_code == 400


def test_admin_no_puede_quitarse_su_propio_rol(test_client, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL
    usuarios = test_client.get('/api/usuarios').json()
    mi_id = next(u["id"] for u in usuarios if u["email"] == TEST_USUARIO_EMAIL)

    resp = test_client.patch(f'/api/usuarios/{mi_id}', json={"rol": "mozo"})
    assert resp.status_code == 400


def test_admin_no_puede_eliminar_su_propia_cuenta(test_client, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL
    usuarios = test_client.get('/api/usuarios').json()
    mi_id = next(u["id"] for u in usuarios if u["email"] == TEST_USUARIO_EMAIL)

    resp = test_client.delete(f'/api/usuarios/{mi_id}')
    assert resp.status_code == 400


def test_un_admin_puede_eliminar_a_otro_admin_distinto(test_client, test_db, test_cliente):
    """No puedes borrarte a vos mismo, pero sí a otro admin del mismo
    restaurante (quedando vos como el admin restante).

    El segundo admin se crea directo en la BD (no vía API): un admin ya
    no puede crear otro admin desde el panel — ver
    test_admin_no_puede_crear_otro_admin — así que para probar el borrado
    hay que insertarlo como si viniera de antes de esa restricción.
    """
    from backend.auth import hash_password
    from backend.models import Usuario

    otro_admin = Usuario(
        id='usr-otro-admin',
        cliente_id=test_cliente.id,
        nombre='Otro Admin',
        email='otroadmin@test-restaurant.com',
        password_hash=hash_password('clave123'),
        rol='admin',
        estado='activo',
    )
    test_db.add(otro_admin)
    test_db.commit()

    resp = test_client.delete(f'/api/usuarios/{otro_admin.id}')
    assert resp.status_code == 204


def test_admin_no_puede_crear_otro_admin(test_client, test_cliente):
    """Cada restaurante tiene un solo admin, dado de alta por el superadmin.
    Un admin no puede crearse un "repuesto" con su mismo nivel de acceso."""
    resp = test_client.post('/api/usuarios', json={
        "nombre": "Otro Admin", "email": "repuesto@test-restaurant.com", "password": "clave123", "rol": "admin",
    })
    assert resp.status_code == 403


def test_admin_no_puede_ascender_a_otro_usuario_a_admin(test_client, test_cliente):
    """Tampoco puede lograrlo en dos pasos: crear un mozo y luego editarlo a admin."""
    creado = test_client.post('/api/usuarios', json={
        "nombre": "Mozo", "email": "mozo-ascenso@test-restaurant.com", "password": "clave123", "rol": "mozo",
    }).json()

    resp = test_client.patch(f'/api/usuarios/{creado["id"]}', json={"rol": "admin"})
    assert resp.status_code == 403


def test_crear_staff_genera_codigo_de_acceso_automaticamente(test_client, test_cliente):
    """El admin no escribe ningún identificador: el backend genera un
    código de 6 dígitos y lo devuelve en la respuesta."""
    resp = test_client.post('/api/usuarios/staff', json={
        "nombre": "Pedro Mozo", "password": "clave123", "rol": "mozo",
    })
    assert resp.status_code == 201
    codigo = resp.json()["celular"]
    assert codigo is not None
    assert len(codigo) == 6
    assert codigo.isdigit()


def test_crear_dos_staff_reciben_codigos_distintos(test_client, test_cliente):
    r1 = test_client.post('/api/usuarios/staff', json={
        "nombre": "Pedro Mozo", "password": "clave123", "rol": "mozo",
    }).json()
    r2 = test_client.post('/api/usuarios/staff', json={
        "nombre": "Ana Cajera", "password": "clave123", "rol": "cajero",
    }).json()
    assert r1["celular"] != r2["celular"]


def test_crear_staff_no_acepta_celular_del_cliente(test_client, test_cliente):
    """El campo celular ya no es parte del payload — si se manda, FastAPI
    lo ignora (extra field) en vez de usarlo como identificador de login."""
    resp = test_client.post('/api/usuarios/staff', json={
        "nombre": "Pedro Mozo", "celular": "999999999", "password": "clave123", "rol": "mozo",
    })
    assert resp.status_code == 201
    assert resp.json()["celular"] != "999999999"


def test_resetear_password_de_usuario_personal(test_client, test_cliente):
    creado = test_client.post('/api/usuarios', json={
        "nombre": "Cajero", "email": "cajero@test-restaurant.com", "password": "vieja12345", "rol": "cajero",
    }).json()

    resp = test_client.patch(f'/api/usuarios/{creado["id"]}/password', json={"nueva_password": "nueva12345"})
    assert resp.status_code == 200


# ============ PERFIL Y CONTRASEÑA ============

def test_usuario_puede_obtener_su_perfil(test_client_real_auth, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL
    token = _login_restaurante(test_client_real_auth)
    test_client_real_auth.headers.update({"Authorization": f"Bearer {token}"})

    resp = test_client_real_auth.get('/api/usuarios/me')
    assert resp.status_code == 200
    data = resp.json()
    assert data["email"] == TEST_USUARIO_EMAIL
    assert data["rol"] == "admin"


def test_usuario_puede_cambiar_su_propia_contraseña(test_client_real_auth, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL, TEST_USUARIO_PASSWORD
    token = _login_restaurante(test_client_real_auth)
    test_client_real_auth.headers.update({"Authorization": f"Bearer {token}"})

    # Cambiar a contraseña nueva
    resp = test_client_real_auth.patch('/api/usuarios/me/password', json={
        "password_actual": TEST_USUARIO_PASSWORD,
        "nueva_password": "nuevaContraseña123"
    })
    assert resp.status_code == 200

    # Verificar que el login falla con contraseña vieja
    test_client_real_auth.headers.clear()
    resp = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL,
        "password": TEST_USUARIO_PASSWORD
    })
    assert resp.status_code == 401

    # Verificar que funciona con contraseña nueva
    resp = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL,
        "password": "nuevaContraseña123"
    })
    assert resp.status_code == 200


def test_usuario_no_puede_cambiar_contraseña_sin_contraseña_actual_correcta(test_client_real_auth, test_cliente):
    token = _login_restaurante(test_client_real_auth)
    test_client_real_auth.headers.update({"Authorization": f"Bearer {token}"})

    resp = test_client_real_auth.patch('/api/usuarios/me/password', json={
        "password_actual": "contraseñaIncorrecta",
        "nueva_password": "otraContraseña123"
    })
    assert resp.status_code == 400
    assert "incorrecta" in resp.json()["detail"].lower()


def test_superadmin_puede_cambiar_su_propia_contraseña(test_client_real_auth_sa):
    from tests.conftest import TEST_SUPERADMIN_PASSWORD
    resp = test_client_real_auth_sa.patch('/api/superadmin/me/password', json={
        "password_actual": TEST_SUPERADMIN_PASSWORD,
        "nueva_password": "nuevoSuperadmin123"
    })
    assert resp.status_code == 200


def test_superadmin_no_puede_cambiar_contraseña_sin_contraseña_actual(test_client_real_auth_sa):
    resp = test_client_real_auth_sa.patch('/api/superadmin/me/password', json={
        "password_actual": "incorrecta",
        "nueva_password": "nueva123"
    })
    assert resp.status_code == 400


def test_superadmin_puede_actualizar_su_nombre(test_client_real_auth_sa):
    resp = test_client_real_auth_sa.patch('/api/superadmin/me', json={
        "nombre": "Juan Pérez Nuevo"
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["nombre"] == "Juan Pérez Nuevo"


def test_superadmin_puede_listar_auditoria(test_client_real_auth_sa):
    """El propio login del superadmin (hecho por el fixture) ya debería
    haber quedado registrado."""
    resp = test_client_real_auth_sa.get('/api/superadmin/auditoria')
    assert resp.status_code == 200
    eventos = resp.json()
    assert any(e["accion"] == "login" and e["entidad"] == "superadmin" for e in eventos)


def test_auditoria_no_accesible_con_token_de_restaurante(test_client_real_auth, test_cliente):
    token = _login_restaurante(test_client_real_auth)
    resp = test_client_real_auth.get('/api/superadmin/auditoria', headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 403


def _login_restaurante(client):
    from tests.conftest import TEST_USUARIO_EMAIL, TEST_USUARIO_PASSWORD
    resp = client.post('/api/auth/login', json={"email": TEST_USUARIO_EMAIL, "password": TEST_USUARIO_PASSWORD})
    return resp.json()["access_token"]


# ============ RATE LIMITING DE LOGIN ============

def test_login_bloquea_tras_varios_intentos_fallidos(test_client_real_auth, test_cliente):
    from tests.conftest import TEST_USUARIO_EMAIL

    for _ in range(5):
        resp = test_client_real_auth.post('/api/auth/login', json={
            "email": TEST_USUARIO_EMAIL, "password": "equivocada",
        })
        assert resp.status_code == 401

    # El 6to intento (aunque la contraseña ahora sí sea correcta) se frena
    # por rate-limit, no por credenciales.
    from tests.conftest import TEST_USUARIO_PASSWORD
    resp = test_client_real_auth.post('/api/auth/login', json={
        "email": TEST_USUARIO_EMAIL, "password": TEST_USUARIO_PASSWORD,
    })
    assert resp.status_code == 429


def test_login_exitoso_no_cuenta_para_el_limite(test_client_real_auth, test_cliente):
    """Varios logins CORRECTOS seguidos (varios mozos entrando desde el
    mismo WiFi al empezar un turno) no deben gastar el límite — solo los
    fallos cuentan."""
    from tests.conftest import TEST_USUARIO_EMAIL, TEST_USUARIO_PASSWORD

    for _ in range(8):
        resp = test_client_real_auth.post('/api/auth/login', json={
            "email": TEST_USUARIO_EMAIL, "password": TEST_USUARIO_PASSWORD,
        })
        assert resp.status_code == 200


def test_login_staff_bloquea_tras_varios_intentos_fallidos(test_client_real_auth, test_cliente):
    for _ in range(5):
        resp = test_client_real_auth.post('/api/auth/login-staff', json={
            "celular": "000000", "password": "equivocada",
        })
        assert resp.status_code == 401

    resp = test_client_real_auth.post('/api/auth/login-staff', json={
        "celular": "000000", "password": "equivocada",
    })
    assert resp.status_code == 429


def test_superadmin_login_bloquea_tras_varios_intentos_fallidos(test_client_real_auth, test_superadmin):
    from tests.conftest import TEST_SUPERADMIN_EMAIL

    for _ in range(5):
        resp = test_client_real_auth.post('/api/superadmin/login', json={
            "email": TEST_SUPERADMIN_EMAIL, "password": "equivocada",
        })
        assert resp.status_code == 401

    resp = test_client_real_auth.post('/api/superadmin/login', json={
        "email": TEST_SUPERADMIN_EMAIL, "password": "equivocada",
    })
    assert resp.status_code == 429


# ============ REGISTRO DE AUDITORÍA ============

def test_login_exitoso_queda_en_el_registro_de_auditoria(test_client_real_auth, test_cliente, test_db):
    from backend.models import AuditLog
    from tests.conftest import TEST_USUARIO_EMAIL

    _login_restaurante(test_client_real_auth)

    evento = test_db.query(AuditLog).filter(AuditLog.accion == "login", AuditLog.actor == TEST_USUARIO_EMAIL).first()
    assert evento is not None
    assert evento.entidad == "usuario"


def test_crear_staff_queda_en_el_registro_de_auditoria(test_client, test_cliente, test_db):
    from backend.models import AuditLog

    creado = test_client.post('/api/usuarios/staff', json={
        "nombre": "Pedro Mozo", "password": "clave123", "rol": "mozo",
    }).json()

    evento = test_db.query(AuditLog).filter(
        AuditLog.accion == "crear_staff", AuditLog.entidad_id == creado["id"],
    ).first()
    assert evento is not None


def test_eliminar_usuario_queda_en_el_registro_de_auditoria(test_client, test_cliente, test_db):
    from backend.models import AuditLog

    creado = test_client.post('/api/usuarios/staff', json={
        "nombre": "Pedro Mozo", "password": "clave123", "rol": "mozo",
    }).json()
    test_client.delete(f'/api/usuarios/{creado["id"]}')

    evento = test_db.query(AuditLog).filter(
        AuditLog.accion == "eliminar_usuario", AuditLog.entidad_id == creado["id"],
    ).first()
    assert evento is not None
